import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from scraper import extract_article, make_level_url
from translations import TRANSLATIONS


def add_note_space(doc, lines):
    for _ in range(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(13)


def add_difficult_words(doc, words):
    if not words:
        return

    line = doc.add_paragraph("_" * 72)
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(4)

    h = doc.add_paragraph()
    run = h.add_run("Difficult words")
    run.bold = True
    run.font.size = Pt(10.5)
    h.paragraph_format.space_after = Pt(3)

    for text in words:
        text = re.sub(r"^Difficult words:\s*", "", text, flags=re.I)
        p = doc.add_paragraph(text)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(5)


def make_safe_title(title):
    title = re.sub(r'[:/\\?*"<>|]', "", title)
    title = re.sub(r"\s+", "_", title).strip("_")
    title = re.sub(r"_+", "_", title)
    return title


def make_output_file(title, publish_dt, output_dir, index=None):
    date_text = publish_dt.date().isoformat()
    safe_title = make_safe_title(title)

    if index is None:
        filename = f"{date_text}_{safe_title}.docx"
    else:
        filename = f"{date_text}_{index:02d}_{safe_title}.docx"

    output_file = output_dir / filename
    if not output_file.exists():
        return output_file

    stem = output_file.stem
    suffix = output_file.suffix
    new_file = output_dir / f"{stem}_new{suffix}"
    if not new_file.exists():
        return new_file

    number = 2
    while True:
        new_file = output_dir / f"{stem}_new{number}{suffix}"
        if not new_file.exists():
            return new_file
        number += 1


def is_article_generated(article, output_dir):
    output_dir = Path(output_dir)
    if not output_dir.exists():
        return False

    date_text = article["publish_dt"].date().isoformat()
    safe_title = make_safe_title(article["title"])
    pattern = re.compile(
        rf"^{re.escape(date_text)}_(?:\d{{2}}_)?{re.escape(safe_title)}(?:_new\d*)?\.docx$",
        re.I,
    )

    return any(pattern.match(path.name) for path in output_dir.glob("*.docx"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.size = Pt(9)

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")

    field_code = OxmlElement("w:instrText")
    field_code.set(qn("xml:space"), "preserve")
    field_code.text = "PAGE"

    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")

    field_text = OxmlElement("w:t")
    field_text.text = "1"

    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")

    run._r.append(field_begin)
    run._r.append(field_code)
    run._r.append(field_separate)
    run._r.append(field_text)
    run._r.append(field_end)


def add_footer_page_number(doc):
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.clear()
    add_page_number(paragraph)


def build_word(title, levels, output_file):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    add_footer_page_number(doc)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.8)

    main_title = doc.add_heading(title, level=0)
    main_title.paragraph_format.space_after = Pt(6)

    for level in [1, 2, 3]:
        if level == 3:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        h = doc.add_heading(f"LEVEL {level}", level=2)
        h.runs[0].bold = True
        h.paragraph_format.space_before = Pt(2)
        h.paragraph_format.space_after = Pt(4)

        paragraphs, difficult_words = levels[level]

        for text in paragraphs:
            p = doc.add_paragraph(text)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.space_after = Pt(5)

        add_difficult_words(doc, difficult_words)

        if level == 1:
            add_note_space(doc, 2)
        elif level == 2:
            add_note_space(doc, 5)
        else:
            add_note_space(doc, 9)

    doc.save(output_file)


def generate_articles(articles, output_dir):
    output_dir.mkdir(exist_ok=True)

    date_counts = defaultdict(int)
    for article in articles:
        date_counts[article["publish_dt"].date().isoformat()] += 1

    date_indexes = defaultdict(int)
    generated_files = []

    for article in articles:
        levels = {}
        article_title = None
        publish_dt = article["publish_dt"]
        date_text = publish_dt.date().isoformat()

        for level in [1, 2, 3]:
            url = make_level_url(article["url"], level)
            title, level_publish_dt, paragraphs, difficult_words = extract_article(url)
            article_title = article_title or title
            publish_dt = level_publish_dt if level == 1 else publish_dt
            levels[level] = (paragraphs, difficult_words)

        if date_counts[date_text] > 1:
            date_indexes[date_text] += 1
            index = date_indexes[date_text]
        else:
            index = None

        output_file = make_output_file(article_title, publish_dt, output_dir, index)
        build_word(article_title, levels, output_file)
        generated_files.append(output_file)

    return generated_files


def get_article_preview(article, language="English"):
    texts = TRANSLATIONS.get(language, TRANSLATIONS["English"])
    preview_parts = [f"{texts['preview_title']}\n{article['title']}"]
    all_difficult_words = []

    for level in [1, 2, 3]:
        url = make_level_url(article["url"], level)
        title, publish_dt, paragraphs, difficult_words = extract_article(url)
        preview_parts.append(f"LEVEL {level}\n" + "\n\n".join(paragraphs))
        all_difficult_words.extend(difficult_words)

    if all_difficult_words:
        clean_words = [
            re.sub(r"^Difficult words:\s*", "", text, flags=re.I)
            for text in all_difficult_words
        ]
        preview_parts.append(f"{texts['difficult_words']}\n" + "\n\n".join(clean_words))
    else:
        preview_parts.append(f"{texts['difficult_words']}\n")

    return "\n\n" + ("-" * 48 + "\n\n").join(preview_parts)
